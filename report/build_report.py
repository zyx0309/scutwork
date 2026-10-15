"""Generate an evidence-based Word/PDF report from the user's actual result tables."""
import csv
import argparse
import hashlib
import html
import json
import math
import shutil
import sys
import urllib.request
from collections import Counter,defaultdict
from datetime import datetime,timezone,timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parent
OUTPUT_ROOT=ROOT
sys.path.insert(0,str(PROJECT/'src'))
sys.path.insert(0,str(ROOT))
from verify_results import verify,csv_rows

FONT_URL='https://raw.githubusercontent.com/notofonts/noto-cjk/main/Sans/Variable/TTF/Subset/NotoSansSC-VF.ttf'
FONT_SHA='d68bafcb48a2707749396aa12bbbd833cb70401f3a9a689fd2902c7e0d295964'
TITLE='基于OpenStreetMap的省级行政区代表城市铁路网络构建及最短几何路径分析'
SUBTITLE='《人工智能》课程大作业 · 第一阶段真实数据实验与质量核验'
COVER_NOTE='数据核验稿：保留已确认站点错配与待复核项'
ATTR='© OpenStreetMap contributors · ODbL 1.0'

def font_file():
    directory=ROOT/'fonts'; directory.mkdir(exist_ok=True)
    target=directory/'NotoSansSC-Regular.ttf'
    if target.exists(): return target
    original=directory/'NotoSansSC-VF.ttf'
    if not original.exists():
        with urllib.request.urlopen(FONT_URL,timeout=90) as response: original.write_bytes(response.read())
    if hashlib.sha256(original.read_bytes()).hexdigest()!=FONT_SHA:
        raise ValueError('Font checksum mismatch; do not disable verification')
    from fontTools.ttLib import TTFont
    from fontTools.varLib.instancer import instantiateVariableFont
    font=instantiateVariableFont(TTFont(original),{'wght':400},inplace=True)
    font.save(target)
    return target

def draw_figures(data,font,diagnostic=None,current=False):
    import os
    os.environ.setdefault('MPLCONFIGDIR',str(PROJECT/'.cache/report-matplotlib'))
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import font_manager,patheffects
    import matplotlib.pyplot as plt
    import numpy as np
    from pyproj import Geod
    font_manager.fontManager.addfont(str(font))
    family=font_manager.FontProperties(fname=font).get_name()
    matplotlib.rcParams.update({'font.family':family,'axes.unicode_minus':False,'svg.fonttype':'path',
                                 'font.size':13,'axes.titlesize':16,'axes.spines.top':False,'axes.spines.right':False})
    target=OUTPUT_ROOT/'figures'; target.mkdir(parents=True,exist_ok=True)
    cities,stations,edges,segs,q=data
    coordinates={r['city']:(float(r['longitude']),float(r['latitude'])) for r in cities}
    offset={'北京':(-25,13),'天津':(13,-3),'太原':(-48,14),'石家庄':(9,-29),
            '南京':(12,13),'合肥':(-35,-12),'上海':(12,-2),'南昌':(12,-7),
            '武汉':(-35,12),'长沙':(-34,-4),'杭州':(10,-15),'广州':(-21,14),
            '香港':(15,-6),'澳门':(7,-28),'海口':(-20,-18),'西宁':(-28,10),'兰州':(12,5),
            '福州':(-16,10),'台北':(14,8),'重庆':(10,-2),'成都':(-22,9)}
    weight_offset={'E099':(0,-13),'E029':(0,4),'E235':(0,8),'E249':(12,-4),
                   'E286':(-14,2),'E181':(10,0),'E162':(10,0),'E059':(0,-9),
                   'E331':(0,10),'E307':(0,-9),'E315':(20,5)}
    paths={}
    def finish(fig,name,caption):
        fig.text(.06,.015,caption+'\n'+ATTR,fontsize=8,color='#55636c')
        fig.savefig(target/(name+'.png'),dpi=190,bbox_inches='tight')
        fig.savefig(target/(name+'.svg'),bbox_inches='tight')
        plt.close(fig); paths[name]=target/(name+'.png')
    # Full graph and a separately labelled minimum spanning forest; no computation edges are removed.
    parent={r['city']:r['city'] for r in cities}
    def find(a):
        while parent[a]!=a: a=parent[a]
        return a
    forest=[]
    for edge in sorted(edges,key=lambda e:(float(e['distance_km']),e['edge_id'])):
        a,b=find(edge['city_a']),find(edge['city_b'])
        if a!=b: parent[a]=b; forest.append(edge)
    with (OUTPUT_ROOT/'mst_display_edges.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=['edge_id','city_a','city_b','distance_km']); writer.writeheader()
        writer.writerows({k:r[k] for k in writer.fieldnames} for r in forest)
    isolated=set(q['city_graph']['isolated_cities'])
    for name,chosen,heading in [('network_full',edges,f'34城市完整连接图：{len(edges)}条边'),
                                ('network_display',forest,f'清晰展示图：几何里程最小生成森林（{len(forest)}条展示边）')]:
        fig,ax=plt.subplots(figsize=(13.2,8.4))
        pending_labels=[]
        for edge in chosen:
            a,b=coordinates[edge['city_a']],coordinates[edge['city_b']]
            focal={edge['city_a'],edge['city_b']}=={'广州','兰州'}
            ax.plot([a[0],b[0]],[a[1],b[1]],color='#d34c41' if focal else '#527f94',
                    lw=1.5 if focal else .9 if name=='network_display' else .45,alpha=.9 if focal else .85 if name=='network_display' else .16)
            if name=='network_display':
                midpoint=((a[0]+b[0])/2,(a[1]+b[1])/2)
                if current:
                    pending_labels.append((edge,midpoint));continue
                ax.annotate(f"{float(edge['distance_km']):.1f}",midpoint,
                        xytext=weight_offset.get(edge['edge_id'],(0,0)),textcoords='offset points',
                        fontsize=10,color='#24536b',ha='center',
                        bbox={'facecolor':'white','alpha':.8,'edgecolor':'none','pad':.5})
        for city,point in coordinates.items():
            color='#b52930' if city in ({'广州','兰州'} if current else {'广州','兰州','福州'}) else '#df9a27' if city in isolated else '#2c6887'
            ax.scatter(*point,s=48 if city in {'广州','兰州'} else 27,c=color,zorder=4,edgecolors='white',linewidths=.5)
            label=city+'*' if city=='福州' and not current else city
            text=ax.annotate(label,point,xytext=offset.get(city,(6,6)),textcoords='offset points',fontsize=13,zorder=5)
            text.set_path_effects([patheffects.withStroke(linewidth=2.8,foreground='white')])
        ax.set_title(heading,pad=15); ax.set_xlabel('经度（OSM城市布局点）'); ax.set_ylabel('纬度（OSM城市布局点）')
        ax.set_xlim(84,130); ax.set_ylim(18,49); ax.grid(alpha=.16); fig.subplots_adjust(bottom=.15)
        if pending_labels:
            # Place current-run distance labels after city labels; old edge IDs changed.
            fig.canvas.draw(); renderer=fig.canvas.get_renderer()
            occupied=[t.get_window_extent(renderer).expanded(1.08,1.16) for t in ax.texts]
            for edge,midpoint in pending_labels:
                annotation=ax.annotate(f"{float(edge['distance_km']):.1f}",midpoint,xytext=(0,0),
                    textcoords='offset points',fontsize=9.5,color='#24536b',ha='center',va='center',
                    bbox={'facecolor':'white','alpha':.88,'edgecolor':'none','pad':.5})
                best=None
                for dy in [0,8,-8,16,-16,24,-24,32,-32,40,-40]:
                    for dx in [0,10,-10,20,-20,30,-30,40,-40]:
                        annotation.set_position((dx,dy));box=annotation.get_window_extent(renderer).expanded(1.1,1.2)
                        overlap=sum(max(0,min(box.x1,b.x1)-max(box.x0,b.x0))*max(0,min(box.y1,b.y1)-max(box.y0,b.y0)) for b in occupied)
                        score=overlap*100+dx*dx+dy*dy
                        if best is None or score<best[0]:best=(score,dx,dy,box)
                _,dx,dy,box=best;annotation.set_position((dx,dy));occupied.append(box)
                if abs(dx)+abs(dy)>12:
                    ax.annotate('',midpoint,xytext=(dx,dy),textcoords='offset points',
                                arrowprops={'arrowstyle':'-','lw':.45,'color':'#6b8c9c'},zorder=1)
        finish(fig,name,(f'仅展示图简化；完整计算图保持{len(edges)}条边。图中数字为km。' if name=='network_display' else '线段仅为抽象城市边，不表示实际轨道走向；红线为广州—兰州城市边。')+
               (' 橙色为孤立城市。' if current else ' 橙色为孤立城市；福州*有已确认交通方式错配。' if diagnostic else ' 橙色为孤立城市；福州*有已确认站名错配。'))
    fig,ax=plt.subplots(figsize=(10.8,4.6))
    labels=['空间匹配待复核','拓扑/关系匹配（其余）','已确认模式错配：福州' if diagnostic else '已确认身份错配：福州','轨道接入未成功','无重铁代表站']
    values=[24,5,1,3,1]
    if current:
        labels=['空间匹配待复核','拓扑/关系匹配','轨道接入未成功','无重铁代表站']
        counts=Counter(r['match_status'] for r in stations)
        values=[counts[k] for k in ['spatial_match_needs_review','topology_matched','station_found_no_reliable_track_match','no_heavy_rail_representative_selected']]
    bars=ax.barh(labels[::-1],values[::-1],color=['#7e8f98','#c2a064','#c7443f','#4f819b','#d9a33a'])
    for bar in bars: ax.text(bar.get_width()+.3,bar.get_y()+bar.get_height()/2,str(int(bar.get_width())),va='center')
    ax.set_xlim(0,27); ax.set_xlabel('城市数量'); ax.set_title('34城市代表站自动匹配结果及审计修正口径')
    fig.tight_layout(rect=(0,.14,1,1)); finish(fig,'mapping_status','自动接入30，其中25个为空间匹配待复核；模式筛查不等于全部股道人工验收。' if current else '6个拓扑/关系匹配标签中含1个已确认错配；自动接入30不等于正确性全部通过。')
    if diagnostic:
        evidence=diagnostic['evidence']['福州']; selected=evidence['selected_track_node']['data']
        train=next(c for c in evidence['name_candidates'] if c['osm_id']==6999105218)
        node_xy={n['osm_id']:(n['lon'],n['lat']) for n in evidence['track_node_coordinates']}
        fig,ax=plt.subplots(figsize=(10.2,6))
        for way in evidence['incident_rail_ways']:
            refs=way['data']['refs']
            if way['source']=='china-latest.osm.pbf' and all(n in node_xy for n in refs):
                geometry=np.asarray([node_xy[n] for n in refs]);ax.plot(geometry[:,0],geometry[:,1],color='#7795a2',lw=1.1)
        ax.scatter(selected['lon'],selected['lat'],s=80,c='#bd3333',zorder=4)
        ax.annotate('本轮国铁轨道锚点\nnode 2451343981' if current else '本轮误选停靠点：subway=yes\nnode 9521844144',
                    (selected['lon'],selected['lat']),xytext=(15,-33),textcoords='offset points',fontsize=11,color='#a32b2b')
        x,y=train['data']['lon'],train['data']['lat'];ax.scatter(x,y,s=90,c='#25816f',marker='s',zorder=4)
        ax.annotate('本轮国铁代表站：train=yes\nnode 6999105218；匹配34.333 m' if current else '国铁站候选：train=yes\nnode 6999105218（未补造锚点）',(x,y),xytext=(12,10),textcoords='offset points',fontsize=11,color='#246c5b')
        if current:
            for way in evidence['incident_rail_ways']:
                refs=way['data']['refs']
                if 2451343981 in refs:
                    line=np.asarray([node_xy[n] for n in refs]);ax.plot(line[:,0],line[:,1],lw=2,color='#256d8a',label=way['data']['tags']['name'])
            ax.legend(loc='upper left',fontsize=11)
        ax.set(xlim=(119.305,119.33),ylim=(26.108,26.123),xlabel='经度',ylabel='纬度',title='福州本轮国铁代表站与实际轨道锚点' if current else '福州同名站交通方式核验：国铁站与滨海快线停靠点')
        ax.ticklabel_format(useOffset=False);ax.grid(alpha=.2);fig.subplots_adjust(bottom=.18)
        finish(fig,'fuzhou_diagnostic','蓝线为锚点所在线路福平线、合福高速线；灰线为其他诊断片段。仅绘制实际坐标，不补造接驳边。' if current else '灰线为诊断包中的实际滨海快线片段；不是国铁福州站完整站场。未绘制两站间虚构连接。')
        fig,axes=plt.subplots(2,3,figsize=(13,9.5))
        for ax,city in zip(axes.flat,['福州','南京','兰州','天津','贵阳','西宁']):
            e=diagnostic['evidence'][city];row=e['selected_station_csv']
            lon,lat=float(row['station_longitude']),float(row['station_latitude'])
            xy={n['osm_id']:(n['lon'],n['lat']) for n in e['track_node_coordinates']}
            for way in e['incident_rail_ways']:
                if all(n in xy for n in way['data']['refs']):
                    line=np.asarray([xy[n] for n in way['data']['refs']]);ax.plot(line[:,0],line[:,1],c='#648396',lw=1)
            ax.scatter(lon,lat,c='#c23b32',s=35,zorder=4)
            ax.scatter(float(row['track_longitude']),float(row['track_latitude']),marker='x',c='#c79725',s=40,zorder=5)
            ax.set(xlim=(lon-.012,lon+.012),ylim=(lat-.009,lat+.009),title=city+'：分量 '+row['component_id'])
            ax.ticklabel_format(useOffset=False);ax.tick_params(labelsize=8);ax.grid(alpha=.18)
        fig.tight_layout(rect=(0,.12,1,1));finish(fig,'focused_track_checks','仅展示诊断所涉轨道片段，非完整站场；红点为CSV定位，金色叉为本轮锚点。不能据图形接近判定连通。')
    else:
        # A real-data error diagnostic: the dashed connector is not a railway.
        fuzhou=coordinates['福州']; taipei=coordinates['台北']; fs=next(r for r in stations if r['city']=='福州')
        wrong=(float(fs['station_longitude']),float(fs['station_latitude']))
        fig,ax=plt.subplots(figsize=(9.5,5.8))
        for point,label,color in [(fuzhou,'福州城市布局点','#295f83'),(wrong,'误匹配：浮洲站（台湾）','#c43f36'),(taipei,'台北城市布局点','#7d8c94')]:
            ax.scatter(*point,s=75,c=color,zorder=3); ax.annotate(label,point,xytext=(9,9 if label!='台北城市布局点' else 28),textcoords='offset points',fontsize=12)
        ax.plot([fuzhou[0],wrong[0]],[fuzhou[1],wrong[1]],'--',c='#bd3e38',lw=1.4)
        ax.text(120.0,25.7,'约246.890 km的身份错配偏移',color='#a62b2b',rotation=-22)
        ax.set(xlim=(118.8,122.7),ylim=(24.6,26.5),xlabel='经度',ylabel='纬度',title='站名消歧案例：福建福州误匹配台湾浮洲')
        ax.grid(alpha=.2); fig.subplots_adjust(bottom=.19)
        finish(fig,'fuzhou_diagnostic','点位来自用户上传cities.csv及city_stations.csv；虚线表示错配偏移，不表示铁路或海峡连接。')
    fig,axes=plt.subplots(1,2,figsize=(10.5,4.8))
    geod=Geod(ellps='WGS84')
    for ax,city in zip(axes,['广州','兰州']):
        r=next(s for s in stations if s['city']==city)
        azimuth,_,distance=geod.inv(float(r['station_longitude']),float(r['station_latitude']),float(r['track_longitude']),float(r['track_latitude']))
        x=distance*math.sin(math.radians(azimuth)); y=distance*math.cos(math.radians(azimuth))
        ax.scatter(0,0,s=95,marker='o',facecolors='none',edgecolors='#2c6a88',label='代表站定位')
        ax.scatter(x,y,s=36,marker='x',color='#c64038',label='选中轨道节点')
        ax.plot([0,x],[0,y],'--',color='#ba8433'); ax.axhline(0,color='#dddddd',lw=.5); ax.axvline(0,color='#dddddd',lw=.5)
        ax.set_xlim(-20,60); ax.set_ylim(-20,60); ax.set_aspect('equal'); ax.grid(alpha=.2)
        ax.set_title(city+f'：偏移{distance:.3f} m'); ax.set_xlabel('相对向东距离（m）'); ax.set_ylabel('相对向北距离（m）')
        ax.legend(fontsize=11,loc='upper left')
    fig.tight_layout(rect=(0,.15,1,1)); finish(fig,'endpoint_offsets','广州定位取OSM站场关系的轨道停靠点；兰州为空间匹配。偏移示意不含周边轨道几何。')
    aggregated=defaultdict(float)
    for r in segs: aggregated[r['line']]+=float(r['distance_km'])
    top=sorted(aggregated.items(),key=lambda r:r[1],reverse=True)[:10]
    fig,ax=plt.subplots(figsize=(10.5,5.9)); ax.barh([x[0] for x in top][::-1],[x[1] for x in top][::-1],color='#427b98')
    for y,(_,value) in enumerate(top[::-1]): ax.text(value+5,y,f'{value:.3f}',va='center',fontsize=12)
    ax.set_xlim(0,690); ax.set_xlabel('同名标签合计几何里程（km）'); ax.set_title('候选路径主要线路标签的里程构成（前10项）')
    fig.tight_layout(rect=(0,.13,1,1)); finish(fig,'line_lengths','据55段CSV原样线路标签合并；统计顺序不表示行车先后，线路名称不证明运营条件。')
    fig,ax=plt.subplots(figsize=(10.8,4.7))
    ax.plot([0]+list(range(1,len(segs)+1)),[0]+[float(r['cumulative_km']) for r in segs],c='#356f91',lw=1.8,marker='.',ms=4)
    ax.set(xlabel='按路径顺序的分段编号',ylabel='累计几何里程（km）',title='55段候选路径累计几何里程')
    ax.grid(alpha=.22); fig.tight_layout(rect=(0,.13,1,1)); finish(fig,'cumulative_distance','横轴为分段编号，不表示时间或地理位置；该曲线不用于推断列车速度。')
    return paths,forest

def build_content(data,audit,images,forest,diagnostic=None):
    cities,stations,edges,segs,q=data; route=q['guangzhou_lanzhou']; blocks=[]
    def h(t,l=1): blocks.append(('h',t,l))
    def p(t):
        for original,replacement in [('10⁻¹³','10^(-13)'),('K₂₇','K27'),('pᵢ₊₁','p_(i+1)'),('pᵢ','p_i'),('p₀','p_0'),('p₁','p_1'),('pₘ','p_m')]:
            t=t.replace(original,replacement)
        blocks.append(('p',t))
    def table(headers,rows,widths=None): blocks.append(('table',headers,rows,widths))
    def image(name,caption): blocks.append(('image',str(images[name]),caption))
    def page(): blocks.append(('page',))
    h('摘　要')
    p('本实验以34个省级行政区代表城市为节点，使用Geofabrik提供的中国与台湾OpenStreetMap数据，构建共享OSM节点连接的无向铁路基础设施图。轨道边权按WGS84椭球上折线全部相邻坐标点的测地距离累计，采用Dijkstra算法计算广州站至兰州站的候选最短几何路径，并将代表站间不经过第三个选定代表站节点的最短通路压缩为城市边。')
    p('用户本地全国数据运行记录显示：底层图含294,713个搜索节点、296,408条边和1,671个连通分量；34城市图含351条边、8个连通分量。程序自动接入30个代表站，广州至兰州候选最短几何里程为2383.541508 km。对上传CSV独立复算得到相同的城市图距离；55段长度之和与总值仅差约9.1×10⁻¹³ km，相邻分段节点连续。')
    p('质量审计同时发现福州被错配到台湾浮洲站，已将该身份匹配标为拒绝；排除后保留29个候选接入，其中24个空间匹配仍待复核，包括兰州站。因此，“自动接入30个”不能等同于“30个代表站正确性已验证”。本文保留原始运行结果与审计状态，将2383.542 km限定为当前筛选规则和站点锚点下的图上候选最短几何里程，不将其表述为官方客运里程或可乘坐列车方案。')
    p('关键词：OpenStreetMap；铁路拓扑；Dijkstra；几何里程；代表站映射；数据质量')
    h('阅读说明')
    p('本报告基于用户上传的8份实际结果文件，不使用人工测试样例作为实验数据。正文区分“程序记录的自检”“本次对导出表的独立核验”与“尚需原始轨道几何或人工检查的事项”。报告没有重新解析全国PBF，也没有补造未上传的实际线路走向图。姓名、学号、专业班级和任课教师均留空待填写。')
    page(); h('1　作业题目与数据建模')
    h('1.1　任务定义与城市口径',2)
    p('题目要求构建34个省级行政区代表城市的铁路网络，以广州到兰州为单方向路径寻优对象。本项目将“34个省会城市及直辖市”统一解释为各省会、自治区首府、直辖市及特别行政区代表城市。节点按用户规定顺序保留，无法定位、没有符合口径的铁路或与其他代表站不连通时，仍保留为孤立节点。')
    p('城市顺序为：'+'、'.join(r['city'] for r in cities)+'。研究不生成34×34距离矩阵，不开展A*或MATLAB对比，也不以未运行实验补足结果。')
    h('1.2　真实数据来源与时间',2)
    p('原始数据由用户在Windows电脑本地提供。中国与台湾文件的PBF复制时间戳均为2026-10-13T20:21:31Z，即北京时间2026年10月14日04:21:31；本次处理完成时间为北京时间2026年10月15日08:52:52。数据复制时间、下载/登记时间和程序完成时间含义不同。')
    sources=q['data_extraction']['sources']
    table(['文件','字节数','PBF复制时间（UTC）','完整解析记录'],[[f'{k}-latest.osm.pbf',str(sources[k]['size']),'2026-10-13 20:21:31','通过'] for k in ['china','taiwan']],[.28,.20,.33,.19])
    p('来源为https://download.geofabrik.de/asia/china-latest.osm.pbf与https://download.geofabrik.de/asia/taiwan-latest.osm.pbf。两文件SHA-256与读取记录见附录。来源清单标记为用户本地文件，未联网核对上游MD5；不能将PBF可解析与本地SHA记录写成已经完成上游发布者校验。OSM数据署名为© OpenStreetMap contributors，许可为ODbL 1.0。')
    h('1.3　底层铁路图与距离定义',2)
    p('采用pyosmium两遍流式扫描，先保留相关铁路way、车站/站场关系及所需节点ID，再扫描坐标和布局节点。两个文件累计扫描268,360,850条原始node记录；这不是去重后的铁路节点数。程序需要的铁路/站点坐标引用为2,335,173个，压缩形状点后的搜索图才是294,713个节点。最终34个城市节点又属于另一层抽象，三者不能混用。')
    p('仅选择正常标签为railway=rail的轨道；地铁、轻轨、有轨电车、废弃/在建及显式工业、军用、纯货运和受限访问线路按代码规则排除。service=yard/siding/spur/crossover不作为任意主干捷径，只保留代表站接入主线所需的最短辅助支路。未知用途不补写为客运，usage=spillway、usage=yard等异常标签也必须在后续清洗中复核。')
    p('轨道仅在共享OSM节点ID上连接。几何交叉、双线临近、车站名称相同都不自动构成轨道连接。平行线路保存各自拓扑，搜索选取某条可行边，不把双线长度累计为一次行程。铁路渡轮和其他非轨道连接不加入图。分析使用无向图，忽略运行方向、时刻表与车辆运营兼容性。')
    p('若轨道折线为P=(p₀,p₁,…,pₘ)，其长度定义为 L(P)=Σ d_WGS84(pᵢ,pᵢ₊₁)/1000，单位km。程序使用pyproj.Geod(ellps="WGS84")求相邻坐标点的椭球测地线距离。所有形状点参与累计，不能对经纬度直接求平面length，也不能只计算两端直线距离。')
    h('1.4　代表车站映射与压缩边',2)
    p('每个城市固定选择一个主要重铁站：如广州站、兰州站、重庆北站、成都东站、台北车站、香港西九龙站。城市中心只用于图形布局，不参与铁路长度计算。映射按站点在轨道上、OSM站场停靠关系、150 m以内最近实际轨道节点的顺序进行。空间匹配仅确定锚点，不创建站房到线路的铁路边，其偏移不计入铁路里程。')
    p('城市边表示两代表站之间存在不经过第三个选定代表站轨道节点的通路，并在此约束下取最短几何长度。允许其他车站及非省会城市。搜索到另一代表站时记录结果且不穿过该节点；普通轨道节点继续扩展。这一定义不等于直达列车，也不等于城市行政区之间没有中途经过。')
    h('2　Dijkstra算法与程序设计')
    h('2.1　最短路径原理',2)
    p('设无向加权图为G=(V,E,w)，几何权重非负。令起点距离dist[s]=0，其余为无穷大，利用优先队列反复取出当前距离最小的候选节点u。对每条邻边(u,v)执行松弛：若dist[u]+w(u,v)<dist[v]，则更新距离及前驱。候选最小节点的距离不能被之后的非负绕行改善，因此逐步得到最短距离。')
    p('二叉堆实现的通常复杂度为O((|V|+|E|)log|V|)，图与前驱存储为O(|V|+|E|)。底层图使用CSR稀疏存储；同一端点对存在平行边时，稀疏搜索权重取最短真实边，同时保留对应OSM轨道标识以恢复路径。')
    h('2.2　伪代码与分层调用',2)
    blocks.append(('code','dist[s] = 0; dist[other] = infinity\nQ = priority_queue((0, s)); predecessor = {}\nwhile Q is not empty:\n    (d, u) = extract_min(Q)\n    if d != dist[u]: continue\n    if city_compression and u is another selected anchor:\n        record_edge_and_path(s, u); continue\n    for (v, weight) in neighbors(u):\n        if d + weight < dist[v]:\n            dist[v] = d + weight\n            predecessor[v] = u\n            push(Q, (dist[v], v))\nrestore track path using predecessor'))
    p('城市压缩搜索由heapq实现，另一个代表站作为停止扩展的屏障；广州至兰州的底层查询使用SciPy的Dijkstra库函数。报告另附独立Python Dijkstra校验脚本，复算上传city_edges.csv中的单方向距离。这是同一算法的独立实现核验，不将其夸大为Dijkstra与A*的多算法实验。')
    h('2.3　实现模块与复现',2)
    table(['模块','职责'],[['extract.py','PBF流式提取、对象类型/ID去重及版本记录'],['stations.py','城市布局、代表站识别与轨道锚点匹配'],['graph.py','折线WGS84长度、轨道图与受约束城市压缩'],['export.py','连续性/长度自检、GeoJSON与CSV导出'],['plotting.py','完整城市图、展示图和车站/路线检查图'],['report/verify_results.py','独立校验上传表格及广州→兰州城市图距离']],[.30,.70])
    p('本地环境为Python 3.12 64位及项目虚拟环境；源码固定依赖包括osmium 4.3.1、pyproj 3.8.0、SciPy 1.17.0、NumPy 2.3.5和Matplotlib 3.10.8。安装依赖需要网络，已有有效PBF的正式处理和绘图默认离线进行。数据、子集SQLite和图缓存分别保存，缓存由输入哈希及策略版本识别。')
    h('3　实验结果及讨论')
    h('3.1　规模、覆盖与代表站结果',2)
    table(['指标','用户运行记录/本次核验结果'],[['城市节点','34，名称和顺序独立核验一致'],['底层搜索图','294,713节点；296,408边；1,671连通分量'],['城市图','351条唯一正权边；8个连通分量'],['自动接入代表站','30/34（88.24%）；不是正确率'],['拓扑/关系匹配标签','6个，其中福州被确认身份错配'],['空间匹配待复核','24个，包含兰州站'],['无可靠轨道锚点','海口、西安、乌鲁木齐'],['无重铁代表站','澳门'],['缺失形状点的way','程序记录0；零长度折线段2条被排除']],[.32,.68])
    p('相关铁路子集中没有记录到跨文件重复对象和版本/内容冲突；该计数不代表对全国所有原始OSM对象进行全量重复审计。中国文件实际node包围框与台湾文件包围框均超出各自头部范围；包围框覆盖与重叠只能作范围提示，不能证明每个区域铁路完整，也不能凭包围框相交断定轨道重复。')
    image('mapping_status','图1　34个代表站匹配状态。原始“拓扑匹配6”中有1个明确身份错配；各项合计34。')
    p('本次共有7个孤立城市：南京、福州、海口、西安、乌鲁木齐、台北、澳门。海口、西安、乌鲁木齐属于所选站没有可靠轨道锚点；南京已经映射但处于不同底层分量；台北位于独立铁路分量；福州的原始孤立状态还受到错误身份映射影响。不能把这些状态统一解释为现实中没有铁路。')
    h('3.2　城市图密度与展示规则',2)
    p('独立根据CSV重新构图，得到27个城市组成的主体分量和7个孤立节点，合计8个分量。主体中每个城市的度均为26，且351=27×26/2，故该主体是完全图K₂₇；全图边密度为351/(34×33/2)=62.57%。')
    p('密图与站点屏障的定义有关：被禁止穿过的是单个选定轨道节点，而不是整个站场或城市区域。线路能够绕过该节点时，仍可能连接任意两端。这不是凭视觉任意补边，也不能仅凭完全图就认定程序直接输出了全城市两两距离；但其邻接含义较宽，必须在报告中解释。')
    image('network_full','图2　按用户cities.csv真实OSM布局坐标重绘的完整34城市图。保留351条边；抽象线段不表示轨道实际线形。福州*为已确认错配城市。')
    image('network_display','图3　仅用于可读展示的几何里程最小生成森林，26条边及其km标注。计算和核验仍使用完整351条边。')
    h('3.3　广州站至兰州站的候选路径',2)
    start,end=route['start_station'],route['end_station']
    table(['项目','起点：广州','终点：兰州'],[['用户选定代表站','广州站','兰州站'],['实际匹配名称',start['station_name'],end['station_name']],['站点OSM对象','relation 8075927','node 269826006'],['轨道锚点ID','8308715569','11698233866'],['轨道经纬度','113.2523777, 23.1529413','103.8486599, 36.0345007'],['匹配方式','站场停靠关系','150 m内轨道节点空间匹配'],['定位偏移','0.000 m（关系停靠点）','38.412 m（待人工复核）']],[.26,.37,.37])
    p('当前底层图的候选最短总长度为2383.541507726979 km，展示为2383.542 km。独立对完整城市边CSV运行Dijkstra得到2383.541507726979 km，与程序报告的底层图距离相同。两个锚点间WGS84测地线下界为1691.584719 km，几何绕行系数约1.409；这属于合理性下界检查，不代替站点身份、实际轨道或运营条件验证。')
    p('“经过的选定代表站城市序列”为广州→兰州，只列出经过的两个选定轨道锚点。它不覆盖中间所有地理城市或普通车站。路径中可识别的站点记录为广州、棠溪、耒阳西、衡山西、白马垅、大托铺、新开铺、捞刀河；报告不根据线路名补造其他停站。')
    image('endpoint_offsets','图4　起终点定位与实际轨道锚点偏移示意。横纵轴为相对米制偏移，虚线不是新增铁路，图中不含周边轨道走向。')
    h('3.4　55段里程与线路构成',2)
    p('路径按可识别停靠节点、线路标签变化分段；未命名线路按OSM way变化分段。上传文件含55段。重新以完整精度求和得到2383.541507726978 km，与总值之差为−9.094947017729282×10⁻¹³ km，可视为双精度累计的舍入差异。相邻段的终点/起点轨道ID全部衔接，首末节点也与起终点锚点一致；累计列逐行复算通过。')
    p('共出现37种原样线路标签，包括未标注线路名。最长的若干合并标签为陇海线589.740 km、京广高速线541.149 km、武西高速线430.007 km、石长线259.584 km与焦柳线144.295 km。未标注线路名合计0.741880 km，占总长度约0.0311%。同名标签合并只是描述构成，不改动原始55段的路径顺序；“京广线”与“京广铁路”保留为两个原始标签。')
    image('line_lengths','图5　按原始标签合并的主要线路几何里程。统计排序不表示行经顺序，不证明相关线路已运营或客车可用。')
    image('cumulative_distance','图6　55段累计几何里程。分段编号是离散路径顺序，不是时间或地理坐标。')
    p('路径还包含“西安东动车走行2线”“三桥西安北动走线”等名称。仅用railway/service/usage等OSM标签筛选不足以证明这些连接适用于实际客运；无向模型也未验证运行方向、轨距、电气化与车辆兼容性。该结果仅是当前基础设施图上的几何寻优结果。')
    h('3.5　已确认缺陷：福州与浮洲错配',2)
    p('city_stations.csv中，福州对应的station_name为“浮洲”，OSM relation ID为10658521，坐标为121.4447595°E、25.0041645°N，来源为taiwan-latest.osm.pbf；而福州城市布局点为119.2918215°E、26.0774954°N。质量记录已经给出约246.890 km的超阈值城市—站点偏移。站名、坐标和来源共同支持这是身份错配，而不是福建福州站的正常站场接入。')
    image('fuzhou_diagnostic','图7　福州—浮洲错配的实际坐标证据。图中的虚线仅是身份错配偏移，不是铁路连接。')
    p('对stations.py的审阅显示，英文Fuzhou别名参与匹配，而福州与浮洲的罗马化名称可能重合；候选排序又优先考虑拓扑关系，这类规则存在跨地域同名站误选风险。现有CSV没有完整候选标签，尚不能确定此次触发的是哪一个名称键或所有候选的排序。站名、位置和来源足以确认身份错配；拓扑关系只能说明对象如何接入轨道，不能单独证明它属于目标城市。')
    p('因此必须把30解释为程序自动接入数，其中至少1个身份已确认有误；其余24个空间匹配也尚未全数人工复核。福州在当前城市CSV中度为0，表级别剔除其错误属性不会改变现有351条边或广州到兰州城市图距离；这不等于已重新运行并验证修正后的全国底层图。')
    p('本报告提供audited_city_stations.csv作为派生审计表：保留原始字段，另设身份审计状态与有效候选锚点列，拒绝福州的浮洲锚点；福州城市节点仍保留为未可靠映射的孤立节点。该表没有补造正确福州站坐标，29表示排除已知错误后的候选数，不表示其余匹配都已人工验收。')
    h('3.6　验证边界与改进建议',2)
    table(['检查层次','本报告可支持的结论'],[['独立核验','34城市顺序；351条正权唯一边；8个城市分量及7个孤立节点；城市图Dijkstra距离；55段求和、累计列、首尾与相邻节点ID衔接'],['引用程序记录','全部城市边底层轨道路径连续及几何重算；最大城市边长度误差0；原始PBF完整解析与扫描规模'],['未独立重算','原始PBF与完整轨道GeoJSON的每个形状点；匹配股道正确性；实际行车方向和运营条件'],['已确认问题','福州→浮洲身份错配；24个空间匹配待审；异常usage标签；稠密K27邻接解释需限定']],[.24,.76])
    p('修正次序应先做站点身份消歧，再复核空间吸附和局部分量。中文主名/规范别名应优先，英文罗马化匹配应同时满足城市/行政区及空间邻近约束；具有强拓扑关系的远程同名站也不能覆盖身份约束。对南京等已接入但孤立节点，应查看局部轨道拓扑；对西安等未接入站，应审查站点对象和停靠关系，不能远距离强行吸附。')
    p('质量记录中的500个“近邻不同分量端点”是程序上限截断的诊断清单，不代表全国恰有500处真实断点，不能据此自动补线。程序还保留7,120条usage未知way、2条usage=spillway与1条usage=yard的main候选，应进一步核对异常标签。保留未知不意味着已确认客运可用。')
    p('如需更有交通解释力的稀疏城市网络，可以在另一次明确版本的模型中把同一代表站站场内的所有真实停靠节点作为屏障集合，或采用有出处的线路邻接定义；不得在本报告中事后删边、补线或修改原始里程来获得预期图形。')
    h('3.7　结论',2)
    p('本实验完成了真实OSM铁路数据到多层图模型的工程流程，并在当前图上得到广州至兰州2383.542 km的候选最短几何结果。独立表级核验支持城市图距离与分段算术的一致性。报告同时明确了福州站身份错配、空间匹配待复核、宽邻接定义与运营约束缺失等限制；计算自洽不能替代数据语义正确性。本版本适合作为带质量审计的实验记录，修正后的代表站识别与实际线路走向仍应补充复核。')
    page(); h('4　程序代码与复现材料')
    p('完整可运行源码随提交包提供；原始大型PBF、虚拟环境和全国图缓存不入包。项目原始分析代码提交为d3a115234a288b57fd46955fb758a882a261f723，仓库为https://github.com/zyx0309/scutwork。报告附加的verify_results.py可对本次上传的8份小型结果重新核验，不需要再次扫描全国PBF。')
    p('复核命令（项目根目录）：.venv\\Scripts\\python.exe report\\verify_results.py report\\input。独立核验在发现站点错误时保留并输出审计项，不把错误标签静默改成正确结果。Word/PDF是相同章节、表格和图表内容的两种排版，分页可能不同。')
    h('参考资料')
    refs=[
      '[1] 用户提供的《人工智能》课程大作业要求及本次OSM第一阶段任务说明。',
      '[2] 用户本次实际运行导出的8份结果文件：cities.csv、city_stations.csv、city_edges.csv、data_quality.json、run_status.json、RESULTS.md、guangzhou_lanzhou_status.json、guangzhou_lanzhou_segments.csv。',
      '[3] Geofabrik. China / Taiwan OSM data extracts. https://download.geofabrik.de/asia.html （具体文件及复制时间见数据来源表）。',
      '[4] OpenStreetMap contributors. Copyright and License. https://www.openstreetmap.org/copyright ; ODbL 1.0: https://opendatacommons.org/licenses/odbl/1-0/ 。',
      '[5] Dijkstra E. W. A Note on Two Problems in Connexion with Graphs. Numerische Mathematik, 1959, 1:269–271. DOI:10.1007/BF01386390.',
      '[6] SciPy Documentation. scipy.sparse.csgraph.dijkstra. https://docs.scipy.org/doc/scipy/reference/generated/scipy.sparse.csgraph.dijkstra.html 。',
      '[7] pyproj Documentation. Geod. https://pyproj4.github.io/pyproj/stable/api/geod.html 。',
      '[8] Pyosmium Documentation. https://docs.osmcode.org/pyosmium/latest/ 。']
    for ref in refs: p(ref)
    page(); h('附录A　34城市原始代表站记录与审计状态')
    p('保持题目规定顺序。分量号为底层轨道图编号，与城市图的8个分量编号不能混用。“关系/节点”仅为自动匹配标签，不代表所有身份与连接正确性已经人工确认。')
    station_table=[]
    for r in stations:
        status={'spatial_match_needs_review':'空间待核','topology_matched':'关系/节点','station_found_no_reliable_track_match':'未可靠接入','no_heavy_rail_representative_selected':'无重铁代表站'}.get(r['match_status'],r['match_status'])
        if r['city']=='福州': status='错配已拒绝'
        station_table.append([r['order'],r['city'],r['station_name'] or '—',
           (r['osm_type'][:1]+'/'+r['osm_id']) if r['osm_id'] else '—',r['track_node_id'] or '—',
           r['match_distance_m'] or '—',r['component_id'] or '—',status])
    table(['序','城市','识别站名','OSM对象','轨道锚点','偏移m','分量','状态'],station_table,[.045,.09,.12,.17,.18,.09,.06,.145])
    page(); h('附录B　广州至兰州55段几何里程')
    p('分段规则为可识别停靠节点或线路标签变化；未命名段按OSM way变化。站名“—”表示该分界没有可靠站名，不能把它当作一个虚构车站。完整起终点轨道节点、OSM way和搜索边ID保留在原始分段CSV，便于追溯。表中km值保留三位小数仅作展示，核验使用原始完整精度。')
    table(['段','线路标签','起点站名','终点站名','本段km','累计km'],
       [[r['segment_id'],r['line'],r['start_station'] or '—',r['end_station'] or '—',f"{float(r['distance_km']):.3f}",f"{float(r['cumulative_km']):.3f}"] for r in segs],
       [.06,.34,.14,.14,.15,.17])
    page(); h('附录C　输入指纹与审计范围')
    p('以下原始PBF哈希来自用户程序来源记录，本次没有重新读取PBF计算这些哈希。用户上传小型结果文件的哈希由本次独立计算，完整值见independent_audit.json和SOURCE_MANIFEST.json。')
    table(['原始文件','程序记录的SHA-256'],[[f'{key}-latest.osm.pbf',sources[key]['sha256'][:32]+'\n'+sources[key]['sha256'][32:]] for key in ['china','taiwan']],[.30,.70])
    p('完整输入数据、图表绘制数据和审计结果作为提交附件保留。城市图及统计图由上传CSV重新生成；本报告未收到原始实际轨道走向图或完整路径GeoJSON，不将抽象城市连线和累计里程曲线当作实际铁路路线图。后续补充路线PNG或GeoJSON后，可加入实际线形图及进一步几何核验。')
    p('数据许可：OSM派生表及地图保留© OpenStreetMap contributors / ODbL 1.0。PDF使用Noto Sans SC字体的子集嵌入，字体许可为SIL Open Font License 1.1，附OFL文本。源代码许可与OSM数据许可分别适用。')
    return diagnosed_blocks(blocks,data,audit,images,diagnostic) if diagnostic else blocks

def diagnosed_blocks(blocks,data,audit,images,diagnostic):
    """Retain the first-run report; explicitly rewrite evidence changed by the second run."""
    completed=datetime.fromisoformat(audit['source_completed_at']).astimezone(timezone(timedelta(hours=8))).strftime('%Y年%m月%d日%H:%M:%S')
    replacements={
      '质量审计同时发现福州被错配到台湾浮洲站':
        '第二轮运行已拒绝台湾浮洲等7个同名/跨地域候选，但缓存诊断发现福州又误选了带subway=yes的本地滨海快线停靠点。本文在派生审计表中拒绝该交通方式映射，保留29个候选接入；其中24个空间匹配仍待复核，包含兰州站。2383.542 km只表示当前图与锚点下的候选最短几何里程，不是官方客运里程或可乘坐列车方案。',
      '本报告基于用户上传的8份实际结果文件':
        '本报告第2版基于同一轮Station_Diagnostics.zip内的8份核心结果、拒绝候选表、缓存诊断JSON和两张实际PNG。运行状态为completed_real_osm，graph_version=5、station_identity_version=2。正文区分程序自检、独立表级核验和局部原始标签证据；没有在云端重读全国PBF或完整轨道GeoJSON。代码的模式修正版为graph_version=6、identity_version=3，已通过测试但尚未全国重跑，不能以预期结果替换本轮实测数据。封面个人信息保留待填。',
      '原始数据由用户在Windows电脑本地提供。':
        '原始数据由用户在Windows电脑本地提供。中国与台湾文件的PBF复制时间均为2026-10-13T20:21:31Z，即北京时间2026年10月14日04:21:31；本轮完成于北京时间'+completed+'。此版本只使用该轮成套输出，不将第一轮数据混入新的独立核验。',
      '仅选择正常标签为railway=rail的轨道':
        '本轮按railway=rail及usage/service/access筛选，排除显式其他轨道类型、废弃/在建和货运/受限线路；站场轨道仅保留代表站接入所需支路。然而本轮未核查停靠关系成员的subway=yes，导致福州代表站交通方式错配，不能声称地铁模式识别已全部正确。未知用途不补写为客运；异常usage标签仍需复核。模式修正版另加显式方式与成员检查，未改变本轮原始输出。',
      '每个城市固定选择一个主要重铁站：':
        '每个城市固定选择一个主要重铁站，城市中心仅用于布局。本轮身份规则优先中文别名并拒绝距城市超过80 km的对象，再按轨道节点、stop_area停靠关系、150 m内最近实际轨道节点接入。站房到轨道的偏移不计入铁路里程。关系中的实际停靠成员仍须核验交通方式；名称和距离合格不能替代这一步。',
      '本实验完成了真实OSM铁路数据到多层图模型的工程流程':
        '本轮当前图上得到广州至兰州2383.542 km候选最短几何结果。独立城市图与55段算术核验一致，已纳入本轮实际路线PNG。审计确认地铁停靠成员误选，并明确南京站场股道、3站锚点缺失及24个空间匹配等限制。新模式门槛已编程测试，尚不能据此宣称全国网络或全部代表站已修复；本报告保留这些已确认问题与验证边界。',
      '完整可运行源码随提交包提供；':
        '完整源码随包提供。生成本轮数据的是站点身份版本2/图版本5，对应代码提交a27a4e833922ef6dd0f138316452ac92c9daffe5；诊断工具后续读取同一份缓存。包内另附模式修正版3/图版本6及STATION_MODE_FIX_GUIDE.md，明确尚未全国重跑。报告不将新源码的预期行为冒充本轮运行结果。原始大型PBF、虚拟环境和全国图缓存不入包。仓库为https://github.com/zyx0309/scutwork。',
      '复核命令（项目根目录）：':
        '表级复核命令（项目根目录）：.venv\\Scripts\\python.exe report\\verify_results.py report\\v2\\input。重建本报告：.venv\\Scripts\\python.exe report\\build_report.py --input report\\v2\\input --output-root report\\v2。无需再扫描全国PBF；缺少完整GeoJSON时不声称逐点几何复算。Word/PDF内容相同，分页可因软件不同。',
      '[2] 用户本次实际运行导出的8份结果文件':
        '[2] 用户提供的本轮Station_Diagnostics.zip：8份核心结果文件、station_candidate_rejections.csv、station_diagnostics.json及车站局部图和广州至兰州路线PNG。局部诊断的原始OSM标签和几何保留来源、对象ID及版本。',
      '完整输入数据、图表绘制数据和审计结果作为提交附件保留。':
        '本轮原始结果和诊断JSON保持原始字节；路线PNG由用户本地程序生成并原样纳入。城市图和局部轨道诊断图另由本轮表格/缓存片段重绘。已收到实际路线图，但未收到完整路径GeoJSON，因此不能凭PNG图像反推每个形状点或宣称独立重算全部轨道长度。',
    }
    output=[];skip=False;seen=set()
    for block in blocks:
        if block[0]=='h' and block[1].startswith('3.5'):
            skip=True
            output.extend([
              ('h','3.5　诊断证据：交通方式漏检与站场局限',2),
              ('p','本轮已拒绝7个候选：太原1个、南京3个、福州2个、南昌1个。浮洲的name:en=Fuzhou、南靖的name:en=Nanjing、南常的英文别名均可追溯，中文/距离门槛已阻止这些对象被选中。7是候选对象数，不是7个失败城市。'),
              ('p','但本轮福州选中relation 19659201“福州火车站”，锚点node 9521844144位于119.3127251°E、26.1123299°N。该节点完整标签明确含railway=stop、subway=yes，关联way 1033520870及1433966622标注为滨海快线。关系本身未标subway，旧程序只检查关系标签，漏掉了成员的模式证据。这是已确认的程序缺陷。'),
              ('image',str(images['fuzhou_diagnostic']),'图8　真实缓存证据中的福州同名站。绿色为train=yes的国铁站候选，红色为本轮误选的subway=yes停靠点；灰线仅为诊断涉及的滨海快线片段。'),
              ('p','同一包中另有国铁福州站node 6999105218，坐标119.3156324°E、26.1168094°N，标签train=yes；stop_area relation 19659200包含该站和若干站台，却没有有效停靠成员可直接接入。识别出正确车站对象不等于已确认它的轨道锚点。本文拒绝本轮地铁锚点，但不补造国铁锚点或通向主图的连接。'),
              ('table',['城市','本地证据','仍不能确定的事项'],[
                 ['南京','国铁站node 267225237在way 384136513“1道”，service=siding；本轮分量7','缺完整站场接入路径，不能凭一条股道判断断连原因'],
                 ['海口','国铁站node 12747885398，train=yes','未找到规则内可靠轨道节点；站点、停靠关系与轨道的联系待核'],
                 ['西安','国铁站node 9620122767，train=yes；另有同名地铁对象','不使用地铁对象替代；国铁锚点仍未可靠匹配'],
                 ['乌鲁木齐','国铁站node 2987146433，另有地铁站node 4676485647','重铁站识别与轨道接入分开；未强行吸附'],
                 ['兰州','国铁站train=yes；所选锚点在陇海线两条way上','偏移38.412 m，股道正确性仍待人工复核'],
              ],[.12,.45,.43]),
              ('image',str(images['focused_track_checks']),'图9　六个重点站的实际轨道证据片段。仅含诊断涉及的way，不是周边完整铁路；轨道空间相近不等于共享节点。'),
              ('p','西宁146.223 m、贵阳123.745 m、天津103.227 m应优先复核。它们分别落在siding、yard、siding标签的股道上；保留必要站场支路可以合理，但距离小于150 m不能证明代表股道正确。兰州的国铁模式和陇海线标签有原始证据支持，仍保留空间匹配待审状态。'),
              ('p','南京与福州的不同分量、海口的海峡非轨道连接等，均不通过自动补线解决。福州分量1109附近的一条15.242 m诊断间隙连接的是分量1110，并非主体分量0；不能将该间隙当成可直接补齐的全国连接。'),
              ('p','模式修正版逐个检查停靠成员并拒绝明确地铁/轻轨/有轨电车等标签，优先明确train=yes的站点证据。可靠车站缺少锚点时保留名称、ID和坐标，保持未接入，不用较弱身份的已连通对象替换。31项测试及小型PBF全流程通过，包含本包真实福州片段回归；尚未在全国PBF上重跑。'),
              ('p','派生audited_city_stations.csv保留全部原始列，拒绝福州的地铁锚点，另列已识别国铁对象6999105218但不赋予新轨道节点。排除这个已确认错误后剩29个候选接入，其中24个空间匹配仍待复核；29不是全部已验证正确的数量。福州城市CSV度为0，故表级拒绝不改变现有351条边及广州至兰州城市图距离，但不等于修正后全国图重建。'),
            ])
            continue
        if skip:
            if block[0]=='h' and block[1].startswith('3.6'): skip=False
            else: continue
        if block[0]=='p':
            text=block[1]
            for prefix,replacement in replacements.items():
                if text.startswith(prefix): text=replacement;seen.add(prefix);break
            text=text.replace('福州的原始孤立状态还受到错误身份映射影响','福州本轮锚点属于被误选的地铁停靠对象')
            output.append(('p',text));continue
        if block[0]=='table':
            rows=[[str(cell).replace('6个，其中福州被确认身份错配','6个，其中福州被确认交通方式错配')
                   .replace('福州→浮洲身份错配','福州停靠点模式错配') for cell in row] for row in block[2]]
            output.append(('table',block[1],rows,block[3]));continue
        if block[0]=='image':
            caption=block[2].replace('明确身份错配','明确交通方式错配')
            if Path(block[1]).stem=='line_lengths':caption=caption.replace('图5','图6')
            if Path(block[1]).stem=='cumulative_distance':caption=caption.replace('图6','图7')
            output.append(('image',block[1],caption))
            if Path(block[1]).stem=='endpoint_offsets':
                output.append(('image',str(images['actual_route']),'图5　本轮Windows真实运行生成的广州至兰州路线PNG，原样纳入。红线沿本地程序的轨道路径绘制；本次没有从图像反推完整坐标或复算其全部几何。'))
            continue
        output.append(block)
    if seen!=set(replacements):raise ValueError('Report rewrite missing expected paragraphs: '+str(set(replacements)-seen))
    return output

def render_word(blocks,path):
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document=Document(); section=document.sections[0]
    section.page_width=Mm(210); section.page_height=Mm(297)
    section.top_margin=Mm(22); section.bottom_margin=Mm(22); section.left_margin=Mm(23); section.right_margin=Mm(23)
    section.header_distance=Mm(9); section.footer_distance=Mm(10)
    section.different_first_page_header_footer=True
    for name in ['Normal','Caption','Heading 1','Heading 2','Heading 3']:
        style=document.styles[name]; style.font.name='Times New Roman'
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体' if name=='Normal' else '微软雅黑')
    normal=document.styles['Normal']; normal.font.size=Pt(10.5); normal.paragraph_format.line_spacing=1.45
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Heading 1',15),('Heading 2',12.5)]:
        style=document.styles[name]; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string('204A64')
        style.paragraph_format.keep_with_next=True; style.paragraph_format.space_before=Pt(12)
    document.styles['Caption'].font.size=Pt(9)
    header=section.header.paragraphs[0]; header.text='人工智能课程大作业　|　OSM铁路网络与最短几何路径'; header.style=document.styles['Caption']
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('第 '); run.font.size=Pt(9)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field); footer.add_run(' 页')
    cover=document.add_paragraph(); cover.paragraph_format.space_before=Pt(40); cover.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=cover.add_run('《人工智能》课程大作业'); run.font.size=Pt(18)
    cover=document.add_paragraph(); cover.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=cover.add_run(TITLE); run.font.size=Pt(23); run.bold=True; run.font.color.rgb=RGBColor.from_string('204A64')
    cover=document.add_paragraph(SUBTITLE); cover.alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph('2026—2027学年第1学期').alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(COVER_NOTE).alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph('')
    for label in ['姓名','学号','专业班级','学院','任课教师','提交日期','学生签名']:
        document.add_paragraph(label+'：____________________________').alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph('真实运行结果日期：2026年10月15日（北京时间）').alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()
    for block in blocks:
        kind=block[0]
        if kind=='h': document.add_heading(block[1],level=block[2])
        elif kind=='p': document.add_paragraph(block[1])
        elif kind=='page': document.add_page_break()
        elif kind=='code':
            p=document.add_paragraph(); p.paragraph_format.line_spacing=1.05
            r=p.add_run(block[1]); r.font.name='Consolas'; r.font.size=Pt(8)
        elif kind=='image':
            p=document.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=1.0; p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(block[1],width=Mm(160))
            p=document.add_paragraph(block[2],style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif kind=='table':
            headers,rows,widths=block[1:]; table=document.add_table(rows=1,cols=len(headers)); table.style='Table Grid'; table.autofit=False
            if widths is None: widths=[1/len(headers)]*len(headers)
            widths=[x/sum(widths) for x in widths]
            for i,(name,width) in enumerate(zip(headers,widths)):
                table.rows[0].cells[i].text=name; table.columns[i].width=Mm(164*width)
            properties=table.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); properties.append(repeat)
            for row in rows:
                cells=table.add_row().cells
                for cell,value in zip(cells,row): cell.text=str(value)
            for index,row in enumerate(table.rows):
                for i,cell in enumerate(row.cells):
                    cell.width=Mm(164*widths[i])
                    if index==0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EFF4'); cell._tc.get_or_add_tcPr().append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing=1.08; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                        for r in p.runs: r.font.size=Pt(8); r.bold=index==0
            document.add_paragraph().paragraph_format.space_after=Pt(1)
    document.core_properties.title=TITLE
    document.core_properties.subject='真实OSM铁路数据；Dijkstra；结果与质量核验'
    document.core_properties.author='课程大作业项目'
    document.save(path)

def render_pdf(blocks,path,font):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image,LongTable,TableStyle,Preformatted,KeepTogether
    pdfmetrics.registerFont(TTFont('NotoSC',str(font)))
    pdfmetrics.registerFontFamily('NotoSC',normal='NotoSC',bold='NotoSC',italic='NotoSC',boldItalic='NotoSC')
    normal=ParagraphStyle('Body',fontName='NotoSC',fontSize=10.3,leading=16.7,spaceAfter=7,wordWrap='CJK')
    head=ParagraphStyle('H1',parent=normal,fontSize=15,leading=21,spaceBefore=12,spaceAfter=9,textColor=colors.HexColor('#204a64'),keepWithNext=True)
    sub=ParagraphStyle('H2',parent=head,fontSize=12.3,leading=18,spaceBefore=10)
    cell=ParagraphStyle('Cell',parent=normal,fontSize=7.7,leading=11,spaceAfter=0)
    caption=ParagraphStyle('Caption',parent=normal,fontSize=8,leading=12,alignment=1,textColor=colors.HexColor('#465f70'),spaceAfter=11)
    centered=ParagraphStyle('Centered',parent=normal,alignment=1,fontSize=13,leading=22)
    cover_title=ParagraphStyle('Title',parent=centered,fontSize=20,leading=30,textColor=colors.HexColor('#204a64'),spaceAfter=14)
    def para(s,style=normal): return Paragraph(html.escape(str(s)).replace('\n','<br/>'),style)
    usable=164*mm; story=[Spacer(1,15*mm),para('《人工智能》课程大作业',centered),Spacer(1,5*mm),para(TITLE,cover_title),
        para(SUBTITLE,centered),para('2026—2027学年第1学期',centered),Spacer(1,5*mm),
        para(COVER_NOTE,centered),Spacer(1,7*mm)]
    for label in ['姓名','学号','专业班级','学院','任课教师','提交日期','学生签名']: story.append(para(label+'：________________________',centered))
    story.extend([Spacer(1,5*mm),para('真实运行结果日期：2026年10月15日（北京时间）',centered),PageBreak()])
    for block in blocks:
        if block[0]=='h': story.append(para(block[1],head if block[2]==1 else sub))
        elif block[0]=='p': story.append(para(block[1]))
        elif block[0]=='page': story.append(PageBreak())
        elif block[0]=='code':
            style=ParagraphStyle('Code',fontName='Courier',fontSize=7.6,leading=11,spaceAfter=9)
            story.append(Preformatted(block[1],style))
        elif block[0]=='image':
            from PIL import Image as PILImage
            with PILImage.open(block[1]) as im: width,height=im.size
            display_width=usable; display_height=display_width*height/width
            maximum_height=(154 if Path(block[1]).stem=='guangzhou_lanzhou_route' else 106)*mm
            if display_height>maximum_height: display_width*=maximum_height/display_height; display_height=maximum_height
            story.append(KeepTogether([Spacer(1,4*mm),Image(block[1],width=display_width,height=display_height),para(block[2],caption)]))
        elif block[0]=='table':
            headers,rows,widths=block[1:]
            if widths is None: widths=[1/len(headers)]*len(headers)
            widths=[x/sum(widths) for x in widths]
            content=[[para(x,cell) for x in row] for row in [headers]+rows]
            table=LongTable(content,colWidths=[usable*x for x in widths],repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7eff4')),
              ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#a8b8c4')),('VALIGN',(0,0),(-1,-1),'TOP'),
              ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
              ('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
            story.extend([table,Spacer(1,8)])
    def footer(canvas,doc):
        canvas.saveState(); canvas.setFont('NotoSC',8); canvas.setFillColor(colors.HexColor('#61717e'))
        if doc.page>1:
            canvas.drawString(23*mm,285*mm,'人工智能课程大作业　|　OSM铁路网络与最短几何路径')
            canvas.drawRightString(187*mm,12*mm,f'第 {doc.page} 页')
        canvas.restoreState()
    doc=SimpleDocTemplate(str(path),pagesize=A4,leftMargin=23*mm,rightMargin=23*mm,topMargin=22*mm,bottomMargin=22*mm,
                          title=TITLE,author='课程大作业项目')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)

def main():
    global OUTPUT_ROOT,SUBTITLE,COVER_NOTE
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=ROOT/'input')
    parser.add_argument('--output-root',type=Path,default=ROOT)
    args=parser.parse_args();OUTPUT_ROOT=args.output_root.resolve();OUTPUT_ROOT.mkdir(parents=True,exist_ok=True)
    folder=args.input.resolve(); audit=verify(folder)
    diagnostic=json.loads((folder/'station_diagnostics.json').read_text(encoding='utf-8-sig')) if (folder/'station_diagnostics.json').exists() else None
    current=(audit['source_graph_version'],audit['source_station_identity_version'])==(6,3)
    if current:
        from reviewed_report import validate_snapshot
        validate_snapshot(folder,args.output_root.resolve()/'REVIEWED_SNAPSHOT.json')
    elif diagnostic and (not audit.get('confirmed_fuzhou_wrong_transport_mode') or audit['source_graph_version']!=5 or audit['source_station_identity_version']!=2):
        raise ValueError('This diagnosed report profile describes the graph-5/identity-2 snapshot. New results require reviewing and updating the report, not reusing old conclusions.')
    if diagnostic:SUBTITLE='《人工智能》课程大作业 · 第二轮真实数据与交通方式审计'
    if current:
        SUBTITLE='《人工智能》课程大作业 · 第三轮真实数据与修正后质量核验'
        COVER_NOTE='真实数据核验稿：已修正模式误选，保留空间匹配待复核项'
    (OUTPUT_ROOT/'independent_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    q=json.loads((folder/'data_quality.json').read_text(encoding='utf-8-sig'))
    data=(csv_rows(folder/'cities.csv'),csv_rows(folder/'city_stations.csv'),csv_rows(folder/'city_edges.csv'),csv_rows(folder/'guangzhou_lanzhou_segments.csv'),q)
    audited=[]
    for original in data[1]:
        row=dict(original); rejected=row['city']=='福州' and (row['station_name']=='浮洲' or audit.get('confirmed_fuzhou_wrong_transport_mode'))
        row['audit_identity_status']=('rejected_wrong_transport_mode' if diagnostic else 'rejected_wrong_station') if rejected else 'not_fully_manually_verified'
        row['candidate_track_node_after_audit']='' if rejected else row['track_node_id']
        row['audit_note']=('福州误选subway=yes停靠点；国铁站6999105218已识别但未补造轨道锚点' if diagnostic else '福州错配台湾浮洲；拒绝该身份映射，保留城市为孤立未定位节点') if rejected else '保留原始自动结果；空间匹配或局部拓扑仍应复核'
        if diagnostic:
            row['identified_train_station_osm_type']='node' if rejected else ''
            row['identified_train_station_osm_id']='6999105218' if rejected else ''
        if current:
            row['audit_identity_status']={'spatial_match_needs_review':'name_mode_screened_spatial_anchor_pending_review',
                'topology_matched':'name_mode_screened_topological_anchor',
                'station_found_no_reliable_track_match':'station_identified_anchor_unresolved',
                'no_heavy_rail_representative_selected':'no_heavy_rail_representative'}[row['match_status']]
            row['audit_note']={'spatial_match_needs_review':'原始名称与模式标签已筛查，偏移独立复算；股道仍待完整站场复核',
                'topology_matched':'保留真实节点或站场关系连接；不表示运营条件全部验证',
                'station_found_no_reliable_track_match':'保留真实站名、OSM对象与坐标；不补造锚点或连接',
                'no_heavy_rail_representative_selected':'保留城市；无本项目口径的重铁代表站'}[row['match_status']]
            if row['city']=='福州':row['audit_note']+='；本轮国铁站train=yes，锚点属于福平线和合福高速线'
        audited.append(row)
    with (OUTPUT_ROOT/'audited_city_stations.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(audited[0])); writer.writeheader(); writer.writerows(audited)
    font=font_file(); images,forest=draw_figures(data,font,diagnostic,current=current)
    if diagnostic:
        for name in ['guangzhou_lanzhou_route.png','station_checks_detail.png']:
            shutil.copy2(folder/'figures'/name,OUTPUT_ROOT/'figures'/name)
        images['actual_route']=OUTPUT_ROOT/'figures/guangzhou_lanzhou_route.png'
    if current:
        from reviewed_report import build_current_content
        blocks=build_current_content(data,audit,images,forest,diagnostic)
    else:blocks=build_content(data,audit,images,forest,diagnostic)
    target=OUTPUT_ROOT/'deliverables'; target.mkdir(exist_ok=True)
    render_word(blocks,target/'AI_Railway_Report_Audited.docx')
    render_pdf(blocks,target/'AI_Railway_Report_Audited.pdf',font)
    md=['# '+TITLE,'\n'+SUBTITLE]
    for b in blocks:
        if b[0]=='h': md.append('\n'+'#'*(b[2]+1)+' '+b[1])
        elif b[0]=='p': md.append('\n'+b[1])
        elif b[0]=='code': md.append('\n```text\n'+b[1]+'\n```')
        elif b[0]=='image': md.append('\n!['+b[2]+'](../figures/'+Path(b[1]).name+')')
        elif b[0]=='table':
            md.append('\n| '+' | '.join(b[1])+' |\n| '+' | '.join(['---']*len(b[1]))+' |')
            md.extend('| '+' | '.join(str(x).replace('\n','<br>') for x in row)+' |' for row in b[2])
    (target/'AI_Railway_Report_Audited.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    print('Word, PDF, Markdown and real-data audit figures generated; input snapshot kept unchanged.')

if __name__=='__main__': main()
